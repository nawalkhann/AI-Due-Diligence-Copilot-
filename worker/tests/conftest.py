from __future__ import annotations

import io
import os
import uuid
from pathlib import Path

import psycopg
import pytest

REPO = Path(__file__).resolve().parents[2]
ADMIN_URL = os.environ.get("DDC_TEST_ADMIN_URL", "postgresql://postgres@localhost:54322/postgres")


# ---------------------------------------------------------------------------
# Sample documents, generated so tests have no binary fixtures to maintain.
# ---------------------------------------------------------------------------

RISK_TEXT = (
    "Our business depends on a small number of large customers. In fiscal 2025 our two "
    "largest customers accounted for 38% of revenue. The loss of either customer could "
    "materially reduce our revenue and operating margin. We also face supply chain risk "
    "because key components are sourced from a single supplier in Taiwan."
)


@pytest.fixture(scope="session")
def sample_pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for n in range(1, 5):
        page = doc.new_page()
        page.insert_text((72, 40), "Acme Robotics Inc. | Annual Report 2025", fontsize=8)
        if n == 1:
            page.insert_text((72, 100), "Item 1A. Risk Factors", fontsize=18)
            page.insert_textbox(pymupdf.Rect(72, 120, 520, 400), RISK_TEXT, fontsize=10)
        elif n == 2:
            page.insert_text((72, 100), "Consolidated Statements of Operations", fontsize=18)
            rows = [["", "2025", "2024"], ["Revenue", "1,250", "1,050"],
                    ["Gross profit", "540", "430"], ["Net income", "120", "85"]]
            y = 140
            for row in rows:
                for col, text in enumerate(row):
                    page.insert_text((72 + col * 150, y), text, fontsize=10)
                y += 20
            for i in range(len(rows) + 1):  # ruling lines so the table is detectable
                page.draw_line((70, 126 + i * 20), (520, 126 + i * 20))
            for x in (70, 220, 370, 520):
                page.draw_line((x, 126), (x, 126 + len(rows) * 20))
        else:
            page.insert_text((72, 100), f"Section {n}", fontsize=18)
            page.insert_textbox(pymupdf.Rect(72, 120, 520, 400),
                                f"Body text for page {n}. Growth in Europe was strong.",
                                fontsize=10)
        page.insert_text((300, 800), str(n), fontsize=8)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture(scope="session")
def sample_html() -> bytes:
    return b"""<html><body>
    <div><b>PART I</b></div>
    <div><b>Item 1A. Risk Factors</b></div>
    <p>""" + RISK_TEXT.encode() + b"""</p>
    <hr style="page-break-after: always">
    <div><b>Item 8. Financial Statements</b></div>
    <table>
      <tr><td></td><td colspan="2">2025</td><td colspan="2">2024</td></tr>
      <tr><td>Revenue</td><td>$</td><td>1,250</td><td>$</td><td>1,050</td></tr>
      <tr><td>Net loss</td><td>$</td><td>(12</td><td>)</td><td>(30)</td></tr>
    </table>
    </body></html>"""


@pytest.fixture(scope="session")
def sample_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    ws.append(["Line item", "FY2024", "FY2025"])
    ws.append(["Revenue", 1050000, 1250000])
    ws.append(["EBITDA", 210000.5, 260000])
    hidden = wb.create_sheet("Scratch")
    hidden.append(["secret"])
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_pptx() -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Investment Highlights"
    slide.placeholders[1].text = "Market leader in warehouse robotics\n40% recurring revenue"
    slide.notes_slide.notes_text_frame.text = "Emphasise net revenue retention of 120%."

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Revenue Growth"
    data = CategoryChartData()
    data.categories = ["2023", "2024", "2025"]
    data.add_series("Revenue ($m)", (800, 1050, 1250))
    slide2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2),
                            Inches(6), Inches(4), data)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_docx() -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.add_heading("Market Overview", level=1)
    doc.add_paragraph("The warehouse automation market is projected to grow 14% annually.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Competitive Landscape", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Competitor", "Share"
    table.cell(1, 0).text, table.cell(1, 1).text = "Locus", "18%"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Database: a fresh database per test session with the real migrations.
# ---------------------------------------------------------------------------

def _can_connect() -> bool:
    try:
        psycopg.connect(ADMIN_URL, connect_timeout=2).close()
        return True
    except psycopg.OperationalError:
        return False


@pytest.fixture(scope="session")
def database_url() -> str:
    if not _can_connect():
        pytest.skip(f"Postgres with pgvector not reachable at {ADMIN_URL}")
    name = f"ddc_test_{uuid.uuid4().hex[:8]}"
    with psycopg.connect(ADMIN_URL, autocommit=True) as admin:
        admin.execute(f'create database "{name}"')
    url = ADMIN_URL.rsplit("/", 1)[0] + f"/{name}"
    with psycopg.connect(url, autocommit=True) as conn:
        conn.execute((REPO / "supabase/tests/local_bootstrap.sql").read_text())
        for migration in sorted((REPO / "supabase/migrations").glob("*.sql")):
            conn.execute(migration.read_text())
    yield url
    with psycopg.connect(ADMIN_URL, autocommit=True) as admin:
        admin.execute(f'drop database "{name}" with (force)')


@pytest.fixture
def db(database_url):
    with psycopg.connect(database_url, autocommit=True) as conn:
        yield conn


def make_user(conn: psycopg.Connection, email: str | None = None) -> str:
    email = email or f"{uuid.uuid4().hex[:8]}@example.com"
    return str(conn.execute("insert into auth.users (email) values (%s) returning id",
                            (email,)).fetchone()[0])


class AsUser:
    """Run statements as an authenticated Supabase user, so RLS applies."""

    def __init__(self, conn: psycopg.Connection, user_id: str) -> None:
        self.conn, self.user_id = conn, user_id

    def __enter__(self):
        self.tx = self.conn.transaction()
        self.tx.__enter__()
        self.conn.execute("set local role authenticated")
        self.conn.execute("select set_config('request.jwt.claim.sub', %s, true)",
                          (self.user_id,))
        return self.conn

    def __exit__(self, *exc):
        return self.tx.__exit__(*exc)
